import io
import json
import uuid
from typing import Any

from httpx import AsyncClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.config import get_settings
from app.core.db import get_sessionmaker, task_session
from app.core.outbox import OutboxEvent
from app.core.storage import MemoryStorage
from app.modules.exhibitions.models import Edition
from app.modules.importexport.models import ImportJob
from app.modules.importexport.rules import cell_to_text, suggest_mapping
from app.modules.importexport.service import process_import
from app.modules.importexport.tasks import run_import_event
from app.modules.visitors.models import VisitorLead
from tests.factories import RoleGrant, actor, make_edition

HEADERS = ["First Name", "Email", "Phone", "Country", "Gender", "Company Name"]
ROWS: list[list[Any]] = [
    ["Tho", "tho@example.test", 84912345678, "VN", "Male", "Khong"],  # 2 ok
    ["Lan", "lan@example.test", None, "VN", "female", None],  # 3 ok
    ["Tho again", "THO@example.test", None, "VN", None, None],  # 4 duplicate in file
    ["Existing", "existing@example.test", None, None, None, None],  # 5 duplicate existing
    ["No contact", None, None, None, None, None],  # 6 invalid: no email/phone
    ["Bad gender", "g@example.test", None, None, "robot", None],  # 7 invalid gender
    [None, None, None, None, None, None],  # blank row, ignored
    ["@Tho", "formula@example.test", None, None, None, None],  # 9 ok
]


def xlsx(headers: list[str], rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.append(headers)
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


async def upload(
    client: AsyncClient, edition: Edition, data: bytes, name: str = "leads.xlsx", **form: Any
) -> Any:
    return await client.post(
        "/api/v1/visitor-leads/imports",
        data={"edition_id": str(edition.id), **form},
        files={
            "file": (
                name,
                data,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


async def run(job_id: str) -> str | None:
    async with task_session() as session:
        status = await process_import(session, uuid.UUID(job_id))
        return status.value if status else None


def test_suggest_mapping_and_cells() -> None:
    mapping = suggest_mapping(["First Name", "E-mail", "Mobile", "Hall No"], {"hall_no": "Hall No"})
    assert mapping == {
        "First Name": "first_name", "E-mail": "email", "Mobile": "phone",
        "Hall No": "custom_fields.hall_no",
    }  # fmt: skip
    assert cell_to_text(84971764.0) == "84971764"
    assert cell_to_text("  ") is None


async def test_import_end_to_end(storage: MemoryStorage) -> None:
    edition = await make_edition()
    _, client = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    r = await client.post(
        "/api/v1/visitor-leads",
        json={"edition_id": str(edition.id), "first_name": "E", "email": "existing@example.test"},
    )
    assert r.status_code == 201

    r = await upload(client, edition, xlsx(HEADERS, ROWS))
    assert r.status_code == 202, r.text
    job_id = r.json()["job_id"]
    assert r.json()["status"] == "QUEUED"

    job = (await client.get(f"/api/v1/visitor-leads/imports/{job_id}")).json()
    assert job["status"] == "QUEUED" and job["total_rows"] is None  # nothing read in the request

    async with get_sessionmaker()() as session:
        event = await session.scalar(
            select(OutboxEvent).where(OutboxEvent.task_name == "importexport.run_import")
        )
    assert event is not None and event.queue == "import"
    assert await run_import_event(event.id) == "COMPLETED_WITH_ERRORS"
    assert await run(job_id) is None  # second run is a no-op

    job = (await client.get(f"/api/v1/visitor-leads/imports/{job_id}")).json()
    assert job["total_rows"] == 7
    assert (
        job["processed_count"],
        job["created_count"],
        job["duplicate_count"],
        job["invalid_count"],
    ) == (7, 3, 2, 2)
    assert job["has_error_report"] is True

    errors = (await client.get(f"/api/v1/visitor-leads/imports/{job_id}/errors")).json()["items"]
    assert [(e["row_number"], e["reason_code"]) for e in errors] == [
        (4, "DUPLICATE_IN_FILE"),
        (5, "DUPLICATE_EXISTING"),
        (6, "INVALID"),
        (7, "INVALID"),
    ]
    assert errors[0]["fields"] == {"email": "Already used in this edition."}
    assert errors[0]["row_preview"]["email"] == "T***@example.test"
    assert "gender" in errors[3]["fields"]

    async with get_sessionmaker()() as session:
        leads = (
            await session.scalars(
                select(VisitorLead).where(VisitorLead.import_job_id == uuid.UUID(job_id))
            )
        ).all()
    assert {lead.first_name for lead in leads} == {"Tho", "Lan", "@Tho"}
    assert all(lead.source == "EXCEL" for lead in leads)
    tho = next(lead for lead in leads if lead.first_name == "Tho")
    assert tho.phone_normalized == "+84912345678" and tho.gender == "MALE"

    r = await client.get(f"/api/v1/visitor-leads/imports/{job_id}/errors.xlsx")
    assert r.status_code == 200
    key = r.json()["url"].split("/")[3].split("?")[0]
    report = load_workbook(io.BytesIO(storage.get("imports", key)), read_only=True)
    report_rows = list(report.worksheets[0].iter_rows(values_only=True))
    assert report_rows[0][:2] == ("Row", "Reason") and len(report_rows) == 5


async def test_reimport_same_file_creates_nothing(storage: MemoryStorage) -> None:
    edition = await make_edition()
    _, client = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    data = xlsx(HEADERS, ROWS[:2])
    first = (await upload(client, edition, data)).json()["job_id"]
    assert await run(first) == "COMPLETED"
    second = (await upload(client, edition, data)).json()["job_id"]
    assert await run(second) == "COMPLETED_WITH_ERRORS"
    job = (await client.get(f"/api/v1/visitor-leads/imports/{second}")).json()
    assert job["created_count"] == 0 and job["duplicate_count"] == 2
    errors = (await client.get(f"/api/v1/visitor-leads/imports/{second}/errors")).json()["items"]
    assert {e["reason_code"] for e in errors} == {"DUPLICATE_EXISTING"}


async def test_resume_from_checkpoint(storage: MemoryStorage) -> None:
    edition = await make_edition()
    _, client = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    job_id = (await upload(client, edition, xlsx(HEADERS, ROWS[:2]))).json()["job_id"]
    async with get_sessionmaker()() as session, session.begin():
        job = await session.get(ImportJob, uuid.UUID(job_id))
        assert job is not None
        job.status = "RUNNING"
        job.checkpoint_row = 2  # row 2 was committed before a worker crash
        job.processed_count = 1
        job.created_count = 1
    assert await run(job_id) == "COMPLETED"
    job_out = (await client.get(f"/api/v1/visitor-leads/imports/{job_id}")).json()
    assert job_out["processed_count"] == 2 and job_out["created_count"] == 2


async def test_manual_mapping_saved_per_edition(storage: MemoryStorage) -> None:
    edition = await make_edition()
    _, client = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    r = await client.post(
        "/api/v1/visitor-custom-fields",
        json={
            "edition_id": str(edition.id),
            "field_key": "hall",
            "label": "Hall",
            "field_type": "text",
        },
    )
    assert r.status_code == 201
    data = xlsx(["Nama", "Surel", "Ruang"], [["Budi", "budi@example.test", "A1"]])
    mapping = {"Nama": "first_name", "Surel": "email", "Ruang": "custom_fields.hall"}

    r = await client.post(
        "/api/v1/visitor-leads/imports/preview",
        data={"edition_id": str(edition.id)},
        files={"file": ("x.xlsx", data, "application/octet-stream")},
    )
    assert r.status_code == 200
    assert r.json()["headers"] == ["Nama", "Surel", "Ruang"]
    assert r.json()["suggested_mapping"] == {}
    assert "custom_fields.hall" in {t["target"] for t in r.json()["targets"]}

    bad = await upload(client, edition, data, column_mapping=json.dumps({"Nama": "password"}))
    assert bad.status_code == 422
    r = await upload(client, edition, data, column_mapping=json.dumps(mapping), save_mapping="true")
    assert await run(r.json()["job_id"]) == "COMPLETED"

    data2 = xlsx(["Nama", "Surel", "Ruang"], [["Sari", "sari@example.test", "B2"]])
    r = await upload(client, edition, data2)  # no mapping sent: saved preset applies
    assert await run(r.json()["job_id"]) == "COMPLETED"
    leads = (
        await client.get(
            "/api/v1/visitor-leads", params={"edition_id": str(edition.id), "q": "sari"}
        )
    ).json()
    lead = (await client.get(f"/api/v1/visitor-leads/{leads['items'][0]['id']}")).json()
    assert lead["custom_fields"] == {"hall": "B2"}

    preview = await client.post(
        "/api/v1/visitor-leads/imports/preview",
        data={"edition_id": str(edition.id)},
        files={"file": ("x.xlsx", data2, "application/octet-stream")},
    )
    assert preview.json()["suggested_mapping"] == mapping


async def test_upload_validation(storage: MemoryStorage, monkeypatch: Any) -> None:
    edition = await make_edition()
    _, client = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    r = await upload(client, edition, b"a,b\n1,2", name="leads.csv")
    assert r.status_code == 422
    r = await upload(client, edition, b"not a zip at all")
    assert r.status_code == 422
    r = await upload(client, edition, xlsx(HEADERS, ROWS), import_type="LEADS_AND_REGISTER")
    assert r.status_code == 422 and "import_type" in r.json()["error"]["fields"]
    monkeypatch.setattr(get_settings(), "import_max_bytes", 100)
    r = await upload(client, edition, xlsx(HEADERS, ROWS))
    assert r.status_code == 422 and r.json()["error"]["code"] == "FILE_TOO_LARGE"
    assert storage.objects == {}


async def test_row_limit_fails_job(storage: MemoryStorage, monkeypatch: Any) -> None:
    edition = await make_edition()
    _, client = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    job_id = (await upload(client, edition, xlsx(HEADERS, ROWS))).json()["job_id"]
    monkeypatch.setattr(get_settings(), "import_max_rows", 3)
    assert await run(job_id) == "FAILED"
    job = (await client.get(f"/api/v1/visitor-leads/imports/{job_id}")).json()
    assert "more than 3 rows" in job["error_summary"]


async def test_import_permissions(storage: MemoryStorage) -> None:
    edition, other = await make_edition(), await make_edition()
    _, agent = await actor(RoleGrant("TELEMARKETING_AGENT", "OWNED", edition.id))
    assert (await upload(agent, edition, xlsx(HEADERS, ROWS))).status_code == 403
    _, team_mgr = await actor(RoleGrant("TELEMARKETING_MANAGER", "OWNED", edition.id))
    assert (await upload(team_mgr, edition, xlsx(HEADERS, ROWS))).status_code == 403
    _, admin = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    job_id = (await upload(admin, edition, xlsx(HEADERS, ROWS))).json()["job_id"]
    _, other_admin = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", other.id))
    assert (await other_admin.get(f"/api/v1/visitor-leads/imports/{job_id}")).status_code == 404
    assert (await upload(other_admin, edition, xlsx(HEADERS, ROWS))).status_code == 404


async def test_template_has_core_and_custom_columns(storage: MemoryStorage) -> None:
    edition = await make_edition()
    _, client = await actor(RoleGrant("EXHIBITION_ADMIN", "EDITION", edition.id))
    await client.post(
        "/api/v1/visitor-custom-fields",
        json={
            "edition_id": str(edition.id),
            "field_key": "hall",
            "label": "Hall",
            "field_type": "text",
        },
    )
    await client.post(
        "/api/v1/visitor-custom-fields",
        json={"edition_id": str(edition.id), "field_key": "internal", "label": "Internal",
              "field_type": "text", "in_import": False},
    )  # fmt: skip
    r = await client.get(
        "/api/v1/visitor-leads/imports/template", params={"edition_id": str(edition.id)}
    )
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    headers = next(wb["Leads"].iter_rows(values_only=True))
    assert headers[:3] == ("First Name", "Email", "Phone")
    assert "Hall" in headers and "Internal" not in headers
