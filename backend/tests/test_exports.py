import io
import uuid
from datetime import UTC, datetime, timedelta
from typing import Any

from httpx import AsyncClient
from openpyxl import load_workbook
from sqlalchemy import select, text

from app.core.config import get_settings
from app.core.db import get_sessionmaker, task_session
from app.core.outbox import OutboxEvent
from app.core.storage import MemoryStorage
from app.modules.identity.models import User
from app.modules.importexport import export_service
from app.modules.importexport.models import ExportJob
from app.modules.importexport.tasks import run_export_event
from tests.factories import RoleGrant, actor
from tests.test_reporting import World, world

EXPORTS = "/api/v1/exports"


async def run_exports() -> list[str | None]:
    async with get_sessionmaker()() as session:
        ids = list(
            await session.scalars(
                select(OutboxEvent.id)
                .where(OutboxEvent.event_type == "export.requested")
                .order_by(OutboxEvent.created_at)
            )
        )
    return [await run_export_event(i) for i in ids]


def workbook(storage: MemoryStorage) -> Any:
    ((_, data),) = [
        (k, v) for k, v in storage.objects.items() if k[0] == get_settings().s3_bucket_exports
    ]
    return load_workbook(io.BytesIO(data), read_only=True)


def rows(ws: Any) -> list[tuple[Any, ...]]:
    return [tuple(r) for r in ws.iter_rows(values_only=True)]


async def request(client: AsyncClient, w: World, **filters: Any) -> Any:
    r = await client.post(f"{EXPORTS}/visitors", json={"edition_id": str(w.edition.id), **filters})
    assert r.status_code == 202, r.text
    assert r.json()["status"] == "QUEUED"
    return r.json()


async def test_visitor_export_end_to_end(storage: MemoryStorage) -> None:
    w = await world()
    async with get_sessionmaker()() as session, session.begin():
        await session.execute(
            text("UPDATE visitor_leads SET first_name = '=HYPERLINK(\"http://x\")' WHERE id = :id"),
            {"id": w.leads[1]},
        )
    job = await request(w.admin, w)
    assert await run_exports() == ["COMPLETED"]
    assert await run_exports() == [None]  # already done, never rebuilt

    r = await w.admin.get(f"{EXPORTS}/{job['id']}")
    body = r.json()
    assert r.headers["cache-control"] == "no-store"
    assert (body["status"], body["row_count"]) == ("COMPLETED", 4)
    assert body["download_url"] == f"{EXPORTS}/{job['id']}/file"
    assert body["file_name"].startswith("visitor-leads-") and body["file_name"].endswith(".xlsx")

    # The file streams through the API (storage is never exposed), for allowed callers only.
    r = await w.admin.get(body["download_url"])
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert r.headers["content-disposition"] == f'attachment; filename="{body["file_name"]}"'
    assert r.headers["cache-control"] == "no-store"
    assert r.content == next(iter(storage.objects.values()))
    assert (await w.team_agent_client.get(body["download_url"])).status_code == 403
    assert (await w.manager.get(body["download_url"])).status_code == 404  # not theirs

    ws = workbook(storage)["Visitor Leads"]
    table = rows(ws)
    header = table[0]
    assert header[:3] == ("Date", "Registration ID", "First Name")
    assert "Type of Buyers" in header and "Lead Owner" in header
    data = {
        r[header.index("Registration ID")]: dict(zip(header, r, strict=True)) for r in table[1:]
    }
    assert len(data) == 4
    values = list(data.values())
    assert any(v["Attend"] == "Attend Sourcing" for v in values)
    assert any(v["Type of Buyers"] == "B-Side" for v in values)
    assert any(v["Valid Lead"] == "Yes" for v in values)
    assert any(v["First Name"] == '\'=HYPERLINK("http://x")' for v in values)
    assert {v["Lead Owner"] for v in values} == {
        w.team_agent.full_name,
        w.other_agent.full_name,
        None,
    }

    async with get_sessionmaker()() as session:
        actions = list(
            await session.scalars(
                text(
                    "SELECT action FROM audit_logs WHERE entity_type = 'export_job' ORDER BY occurred_at"
                )
            )
        )
    assert actions == ["requested", "completed", "download"]


async def test_export_follows_scope_and_filters(storage: MemoryStorage) -> None:
    w = await world()
    await request(w.manager, w)
    await run_exports()
    ws = workbook(storage)["Visitor Leads"]
    assert len(rows(ws)) - 1 == 2  # only the team agent's leads

    storage.objects.clear()
    await request(w.admin, w, source="EXCEL")
    await run_exports()
    assert len(rows(workbook(storage)["Visitor Leads"])) - 1 == 1


async def test_export_visibility(storage: MemoryStorage) -> None:
    w = await world()
    job = await request(w.manager, w)
    assert (await w.admin.get(f"{EXPORTS}/{job['id']}")).status_code == 200
    _, other_manager = await actor(RoleGrant("TELEMARKETING_MANAGER", "OWNED", w.edition.id))
    assert (await other_manager.get(f"{EXPORTS}/{job['id']}")).status_code == 404
    assert (await w.team_agent_client.get(f"{EXPORTS}/{job['id']}")).status_code == 403
    assert (await w.admin.get(f"{EXPORTS}/{uuid.uuid4()}")).status_code == 404


async def test_requester_who_lost_access_gets_failed_export(storage: MemoryStorage) -> None:
    w = await world()
    job = await request(w.manager, w)
    async with get_sessionmaker()() as session, session.begin():
        await session.execute(
            text("UPDATE users SET is_active = false WHERE id = :id"),
            {"id": (await _requester(job["id"]))},
        )
    assert await run_exports() == ["FAILED"]
    async with get_sessionmaker()() as session:
        row = await session.get(ExportJob, uuid.UUID(job["id"]))
    assert row is not None and "no longer active" in (row.error_summary or "")
    assert storage.objects == {}


async def _requester(job_id: str) -> uuid.UUID:
    async with get_sessionmaker()() as session:
        row = await session.get(ExportJob, uuid.UUID(job_id))
        assert row is not None
        user = await session.get(User, row.requested_by)
        assert user is not None
        return user.id


async def test_report_export(storage: MemoryStorage) -> None:
    w = await world()
    r = await w.admin.post(
        f"{EXPORTS}/reports", json={"edition_id": str(w.edition.id), "report": "LEADS"}
    )
    assert r.status_code == 202, r.text
    assert await run_exports() == ["COMPLETED"]
    wb = workbook(storage)
    summary = dict(rows(wb["Summary"])[1:])
    assert (summary["Total"], summary["Unassigned"]) == (4, 1)
    assert rows(wb["By source"])[0] == ("Key", "Count")

    denied = await w.team_agent_client.post(
        f"{EXPORTS}/reports", json={"edition_id": str(w.edition.id), "report": "LEADS"}
    )
    assert denied.status_code == 403
    edition_only = await w.manager.post(
        f"{EXPORTS}/reports", json={"edition_id": str(w.edition.id), "report": "EMAILS"}
    )
    assert edition_only.status_code == 403


async def test_cleanup_expires_files_and_fails_stuck_jobs(storage: MemoryStorage) -> None:
    w = await world()
    done = await request(w.admin, w)
    await run_exports()
    stuck = await request(w.admin, w)
    past = datetime.now(UTC) - timedelta(hours=2)
    async with get_sessionmaker()() as session, session.begin():
        await session.execute(
            text("UPDATE export_jobs SET expires_at = :t WHERE id = :id"),
            {"t": past, "id": done["id"]},
        )
        await session.execute(
            text("UPDATE export_jobs SET status = 'RUNNING', started_at = :t WHERE id = :id"),
            {"t": past, "id": stuck["id"]},
        )
    async with task_session() as session:
        assert await export_service.cleanup(session) == 2
    assert storage.objects == {}
    statuses = {j["id"]: (await w.admin.get(f"{EXPORTS}/{j['id']}")).json() for j in (done, stuck)}
    assert statuses[done["id"]]["status"] == "EXPIRED"
    assert statuses[done["id"]]["download_url"] is None
    assert (await w.admin.get(f"{EXPORTS}/{done['id']}/file")).status_code == 404
    assert statuses[stuck["id"]]["status"] == "FAILED"


async def test_list_exports_own_or_edition_wide(storage: MemoryStorage) -> None:
    w = await world()
    mine = await request(w.manager, w)
    theirs = await request(w.admin, w)
    params = {"edition_id": str(w.edition.id)}
    manager_view = (await w.manager.get(EXPORTS, params=params)).json()
    assert [j["id"] for j in manager_view["items"]] == [mine["id"]]
    admin_view = (await w.admin.get(EXPORTS, params=params)).json()
    assert {j["id"] for j in admin_view["items"]} == {mine["id"], theirs["id"]}
    assert all(j["download_url"] is None for j in admin_view["items"])
    assert (await w.team_agent_client.get(EXPORTS, params=params)).status_code == 403
