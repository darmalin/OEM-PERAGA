"""Excel helpers. Call these from workers (or a threadpool for small files), never inline
in an async request handler."""

import io
from collections.abc import Iterable, Iterator
from typing import Any

from openpyxl import Workbook, load_workbook

_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def safe_cell(value: Any) -> Any:
    """Neutralize formula injection in exported text (BACKEND_RULES/jobs-integrations.md)."""
    if isinstance(value, str) and value.startswith(_FORMULA_PREFIXES):
        return "'" + value
    return value


def write_workbook(sheets: Iterable[tuple[str, list[str], Iterable[Iterable[Any]]]]) -> bytes:
    """Write-only workbook; rows are streamed, never held as cell objects."""
    wb = Workbook(write_only=True)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title=title[:31])
        ws.append(headers)
        for row in rows:
            ws.append([safe_cell(v) for v in row])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


class WorkbookWriter:
    """Write-only workbook filled sheet by sheet, so an async caller can append in batches."""

    def __init__(self) -> None:
        self._wb = Workbook(write_only=True)
        self._ws: Any = None

    def sheet(self, title: str, headers: list[str]) -> None:
        self._ws = self._wb.create_sheet(title=title[:31])
        self._ws.append(headers)

    def append(self, row: Iterable[Any]) -> None:
        self._ws.append([safe_cell(v) for v in row])

    def to_bytes(self) -> bytes:
        buf = io.BytesIO()
        self._wb.save(buf)
        return buf.getvalue()


def iter_rows(data: bytes) -> Iterator[tuple[int, tuple[Any, ...]]]:
    """(spreadsheet row number, values) for the first sheet, streaming in read-only mode."""
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        for number, values in enumerate(ws.iter_rows(values_only=True), start=1):
            yield number, tuple(values)
    finally:
        wb.close()


def is_xlsx(data: bytes) -> bool:
    """xlsx is a zip container."""
    return data[:4] == b"PK\x03\x04"
